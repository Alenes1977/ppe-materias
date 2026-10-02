from pathlib import Path
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
import json

root = Path.cwd()
path = root / 'gestion-aplicada-contraste-fuentes.docx'
buffer = BytesIO()
with ZipFile(path) as z, ZipFile(buffer, 'w', ZIP_DEFLATED) as out:
    for info in z.infolist():
        out.writestr(info.filename.replace('\\', '/'), z.read(info))
buffer.seek(0)
doc = Document(buffer)
p = list(doc.paragraphs)
t = list(doc.tables)

def replace(paragraph, text, style=None):
    paragraph.text = text
    if style:
        paragraph.style = style

def table_to_paragraph(table, text):
    element = OxmlElement('w:p')
    table._element.addprevious(element)
    from docx.text.paragraph import Paragraph
    para = Paragraph(element, table._parent)
    para.text = text
    table._element.getparent().remove(table._element)
    return para

replace(p[0], 'Contraste del plan de estudios de Gestión Aplicada', 'Title')
replace(p[1], 'Informe para revisión por los responsables del título', 'Subtitle')
replace(p[2], 'Actualización del 2 de octubre de 2026 con el anexo vigente')
table_to_paragraph(t[0], 'La revisión toma como referencia canónica el Verifica vigente del título, ID 2503801, de 26 de marzo de 2026. El anexo 4.1 recibido permite completar los 48 pares de ponderaciones del JSON. Se mantienen las incidencias detectadas en las cabeceras web y en la matriz de competencias, y se señalan aspectos del propio anexo que requieren aclaración.')
replace(p[7], 'Los ECTS de las asignaturas de la web y del JSON coinciden con la distribución temporal por materia del Verifica y del anexo 4.1. También coinciden las siete asignaturas básicas que el anexo enumera con nombre, carácter, curso y semestre. La correspondencia del anexo se acredita por el CSV 966597729780096717618167, idéntico al citado en la p. 23 del PDF principal.')
replace(p[8], 'El anexo tampoco enumera todas las asignaturas concretas. Los nombres restantes del JSON proceden de la web y son compatibles con los repartos temporales canónicos. Debe confirmarse con los responsables el detalle que el propio Verifica no explicita; la falta de acceso al anexo ya está resuelta.')
t[2].cell(1,2).text = 'El anexo confirma las asignaturas básicas de Gestión Empresarial 1 y 2 y Contabilidad y Finanzas 1 (pp. 6 y 11). Las restantes no se enumeran individualmente.'
t[2].cell(2,2).text = 'El anexo confirma Comunicación Empresarial 1 y 2 (p. 14) y Segundo y Tercer Idioma Moderno 1 como anuales (p. 16). El reparto por materia es compatible.'
t[2].cell(3,2).text = 'El anexo confirma 6 ECTS en el semestre 6 y 6 en el 7 (p. 9), pero no enumera Derecho 1 y 2; sus nombres proceden de la web.'
t[2].cell(4,2).text = 'El anexo confirma los repartos de estas materias (pp. 18–24). Formación Complementaria incluye intercambio y también optativos de la propia Universidad sin una oferta individualizada.'
replace(p[12], p[12].text + ' El anexo confirma RA12 y RA13 en Empresa y Entorno y CE8 como RA27 (p. 7); también confirma RA16, sin RA20, en Comunicación en las Organizaciones (p. 14). Estas diferencias de la matriz quedan acreditadas por el Verifica vigente.')
replace(p[13], 'Evaluación acreditada por el anexo', 'Heading 1')
replace(p[14], 'Ponderaciones mínimas y máximas', 'Heading 2')
table_to_paragraph(t[4], 'Se incorporaron al JSON los 48 pares de mínimos y máximos de las nueve materias, transcritos de las tablas del anexo y comprobados visualmente. La retirada anterior queda como antecedente: la carencia documental que la motivó está resuelta. Los resultados de aprendizaje, las actividades formativas y los sistemas de evaluación ya coincidían y se conservaron.')
replace(p[15], 'Son horquillas por materia. Los extremos no tienen que sumar 100; los pesos efectivos de cada evaluación sí deben hacerlo y respetar los límites aplicables. No se han generado porcentajes específicos por asignatura. En Formación Complementaria, el anexo advierte que los datos de intercambio son orientativos y dependen de la universidad de destino (p. 20).')
replace(p[16], 'Fuentes de los porcentajes', 'Heading 2')
t[5].cell(1,1).text = 'El PDF principal describe los sistemas (p. 18); el anexo recibido detalla las horquillas de las nueve materias. Su CSV coincide con la referencia canónica.'
t[5].cell(1,2).text = 'Fuente utilizada para completar los 48 pares de límites. Revisión directa del anexo recibido el 2 de octubre de 2026.'
t[5].cell(2,2).text = 'Antecedente de la búsqueda. Los límites actuales se transcriben del anexo vigente; la coincidencia histórica no sustituye esa acreditación.'
replace(p[18], 'Confirmar con los responsables el detalle de asignaturas que el Verifica no enumera y las diferencias de asignación de resultados en la matriz de competencias.')
replace(p[22], 'Anexo 4.1 recibido, 25 páginas, CSV 966597729780096717618167. Copia incluida en el paquete de revisión.')
replace(p[27], 'Web, guías y Google Sheets consultados el 25 de septiembre de 2026. Anexo recibido y contrastado el 2 de octubre de 2026. Las referencias al anexo indican posiciones dentro del PDF; su página 13 muestra «3» en el pie.')

# Add the verified ranges directly before the source summary.
heading = p[16].insert_paragraph_before('Horquillas transcritas por materia', style='Heading 2')
parsed = json.loads((root / 'tmp/pdfs/annex-parsed.json').read_text(encoding='utf-8'))
pages = ['9','11','13','16','17–18','19','20–21','23','25']
table = doc.add_table(rows=1, cols=3)
table.rows[0].cells[0].text = 'Materia'
table.rows[0].cells[1].text = 'Mínimo–máximo por sistema'
table.rows[0].cells[2].text = 'Página PDF'
for source, page in zip(parsed, pages):
    cells = table.add_row().cells
    cells[0].text = source['name']
    cells[1].text = '\n'.join(f'{code}: {low}–{high} %' for code, low, high in source['evaluation'])
    cells[2].text = page
heading._p.addnext(table._element)

# Source inconsistencies matter to the reviewing recipients but do not justify inferred JSON edits.
anchor = p[17]
anchor.insert_paragraph_before('Aspectos del anexo que requieren aclaración', style='Heading 1')
items = [
    'Core Curriculum: el texto general habla de dos asignaturas obligatorias y dos optativas de 3 ECTS (p. 4), mientras que la ficha específica declara 18 ECTS obligatorios (p. 24), como el PDF principal. El JSON mantiene la ficha específica; se solicita revisar la descripción general.',
    'Horas: Empresa y Entorno suma 1.971 horas para 78 ECTS (p. 9) y Comunicación en las Organizaciones, 681 para 27 ECTS (p. 15). A 25 horas por ECTS, razón que arrojan las restantes fichas, serían 1.950 y 675. Se solicita confirmar el criterio de cómputo.',
    'Prácticas: AF5 figura con 550 horas y 0 % de presencialidad (p. 19), frente al 100 % de AF5 en otras fichas. Se solicita aclarar el criterio para las actividades en empresas.',
    'El esquema actual del JSON no registra horas, presencialidad, idiomas de impartición ni contenidos. El anexo conserva ese detalle. Tampoco se inventa una asignación de metodologías MD por materia que el anexo no proporciona.'
]
for text in items:
    anchor.insert_paragraph_before(text, style='List Bullet')

for idx in [4,6,10,17,20]:
    p[idx].style = 'Heading 1'
for style_name in ['Normal','List Bullet']:
    style = doc.styles[style_name]
    style.font.name = 'Calibri'
    style.font.size = Pt(10.5)
    style.font.color.rgb = RGBColor(0,0,0)
    style.paragraph_format.space_after = Pt(7)
    style.paragraph_format.line_spacing = 1.12
for style_name, size in [('Title',22),('Subtitle',11),('Heading 1',14),('Heading 2',12)]:
    style = doc.styles[style_name]
    style.font.name = 'Calibri'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0,0,0)
    style.paragraph_format.keep_with_next = True
    for para in doc.paragraphs:
        if para.style.name == style_name:
            for run in para.runs:
                run.font.color.rgb = RGBColor(0,0,0)
                run.font.size = Pt(size)
for section in doc.sections:
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = section.right_margin = Cm(2)
    section.top_margin = section.bottom_margin = Cm(1.8)
for tab in doc.tables:
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    tab.autofit = False
    ncols = len(tab.columns)
    widths = [3.4,11.8,1.8] if tab is table else ([4,3.25,3.25,3.25,3.25] if ncols == 5 else [3.6,7.5,5.9])
    for col, width in zip(tab.columns,widths): col.width = Cm(width)
    borders = OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        child=OxmlElement('w:'+edge)
        child.set(qn('w:val'),'single');child.set(qn('w:sz'),'4');child.set(qn('w:color'),'D9D9D9')
        borders.append(child)
    props=tab._element.tblPr
    for old in props.findall(qn('w:tblBorders')): props.remove(old)
    props.append(borders)
    for ri,row in enumerate(tab.rows):
        if ri == 0:
            repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
        for ci,cell in enumerate(row.cells):
            if ci < len(widths): cell.width = Cm(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr()
            for old in pr.findall(qn('w:shd')): pr.remove(old)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7EDF3' if ri==0 else ('FFFFFF' if ri%2 else 'F7F8FA'));pr.append(shade)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                child=OxmlElement('w:'+edge);child.set(qn('w:w'),'100');child.set(qn('w:type'),'dxa');margins.append(child)
            pr.append(margins)
            for para in cell.paragraphs:
                para.paragraph_format.keep_with_next = False
                para.paragraph_format.space_after = Pt(3)
                para.paragraph_format.line_spacing = 1.05
                if tab is table and ci==2: para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    run.font.name='Calibri';run.font.size=Pt(9.5);run.font.color.rgb=RGBColor(0,0,0)
                    if ri==0: run.bold=True
doc.core_properties.title='Contraste del plan de estudios de Gestión Aplicada'
doc.core_properties.subject='Actualización con el anexo vigente y revisión de fuentes'
doc.save(path)
print('Informe Word actualizado; ZIP OOXML normalizado y', len(doc.tables), 'tablas.')
